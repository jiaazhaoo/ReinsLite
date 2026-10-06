"""C0: lint table headers and Markdown prose against glossary.toml.

Code is not linted yet (Python's `pass`, loop variables named `plan`, ... would drown the signal); table headers are
where an ambiguous word reaches another person, and docs are where meanings get fixed.
"""
from __future__ import annotations

import csv
import re
import tomllib
from dataclasses import dataclass
from pathlib import Path

GLOSSARY = Path(__file__).with_name("glossary.toml")


@dataclass
class Hit:
    where: str
    word: str
    term: str
    hint: str


def load(path: Path = GLOSSARY) -> list[dict]:
    terms = tomllib.loads(path.read_text(encoding="utf-8"))["term"]
    seen = set()
    for t in terms:
        if t["term"] in seen:
            raise ValueError(f"glossary: {t['term']!r} defined twice")
        seen.add(t["term"])
    return terms


def words(identifier: str) -> list[str]:
    """'TraceConfidence' / 'trace-confidence' / 'trace confidence' -> ['trace', 'confidence']."""
    s = re.sub(r"([a-z0-9])([A-Z])", r"\1_\2", identifier)
    return [w for w in re.split(r"[^A-Za-z0-9]+", s.lower()) if w]


def check_identifier(identifier: str, terms: list[dict]) -> list[tuple[str, dict]]:
    ws = words(identifier)
    if not ws:
        return []
    whole = "_".join(ws)
    # a term may forbid its own bare form ("confidence" must always carry a stage prefix)
    canonical = {t["term"] for t in terms if not t.get("retired") and f"={t['term']}" not in t.get("forbidden", [])}
    if whole in canonical:
        return []
    out = []
    for t in terms:
        if t.get("retired"):
            continue
        for f in t.get("forbidden", []):
            if f.startswith("="):
                if whole == "_".join(words(f[1:])):
                    out.append((f[1:], t))
            else:
                fw = words(f)
                if any(ws[i:i + len(fw)] == fw for i in range(len(ws) - len(fw) + 1)):
                    out.append((f, t))
    once = {}
    for f, t in out:                     # "ground_truth" hits both "truth" and "=ground_truth": report the term once
        once.setdefault(t["term"], (f, t))
    return list(once.values())


def _hint(t: dict) -> str:
    v = f" (values: {', '.join(t['values'])})" if t.get("values") else ""
    return f"use {t['term']}{v}: {t['definition']}"


def lint_headers(path: Path, terms: list[dict]) -> list[Hit]:
    if path.suffix.lower() == ".xlsx":
        try:
            import openpyxl
        except ImportError:
            return [Hit(str(path), "", "", "openpyxl not installed; xlsx headers not checked")]
        wb = openpyxl.load_workbook(path, read_only=True)
        sheets = [(f"{path}[{ws.title}]", [c for c in next(ws.iter_rows(max_row=1, values_only=True), ()) if c])
                  for ws in wb.worksheets]
    else:
        with path.open(encoding="utf-8-sig", newline="") as f:
            header = next(csv.reader(f, delimiter="\t" if path.suffix.lower() == ".tsv" else ","), [])
        sheets = [(str(path), header)]
    return [Hit(f"{where} column {col!r}", w, t["term"], _hint(t))
            for where, cols in sheets for col in cols for w, t in check_identifier(str(col), terms)]


def lint_markdown(path: Path, terms: list[dict]) -> list[Hit]:
    """Prose only: fenced blocks and `inline code` are references to names, not uses of words."""
    hits, fenced = [], False
    for n, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
        if line.lstrip().startswith("```"):
            fenced = not fenced
            continue
        if fenced:
            continue
        prose = re.sub(r"`[^`]*`", " ", line)
        for tok in re.findall(r"[A-Za-z][A-Za-z0-9_]*", prose):
            for w, t in check_identifier(tok, terms):
                hits.append(Hit(f"{path}:{n}", tok, t["term"], _hint(t)))
    return hits


def lint(paths: list[Path], terms: list[dict] | None = None) -> list[Hit]:
    terms = terms or load()
    hits = []
    for p in paths:
        files = sorted(p.rglob("*")) if p.is_dir() else [p]
        for f in files:
            suf = f.suffix.lower()
            if suf in (".csv", ".tsv", ".xlsx"):
                hits += lint_headers(f, terms)
            elif suf == ".md":
                hits += lint_markdown(f, terms)
    return hits
