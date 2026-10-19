"""The production report as a workbook: the same numbers as the board's production page (reins/production.py), for
people who open a spreadsheet, not a web page.

    reins report xlsx [--out FILE] [--lang en|zh] [--hide site,spend]
    reins report publish [--lang ...] [--hide ...]      write it, then copy it to [report] remote with rclone

Sheets: Summary, Work packages, Batches (one row per production / rework batch, comparable), Development, Daily.
--hide site   the first scope part (e.g. the council) becomes "Client A", "Client B", ...
--hide spend  no money anywhere
--lang en     no CJK text at all: a batch alias with CJK is replaced by its batch id, change descriptions are left out

$REINS_HOME/config.toml:
    [report]
    remote = "onedrive:Reins/production.xlsx"     # any rclone destination; the file is replaced, its share link stays
    lang = "en"
    hide = []
    out = ""                                      # default $REINS_HOME/reports/production-<lang>.xlsx
"""
from __future__ import annotations

import re
import subprocess
from contextlib import closing
from pathlib import Path

from .store import ReinsError, config, connect, home

CJK = re.compile(r"[　-鿿＀-￯]")

WORDS = {
    "en": {
        "Summary": "Summary", "Work packages": "Work packages", "Batches": "Batches", "Development": "Development", "Daily": "Daily",
        "title": "Production report", "updated": "Updated", "cost": "Cost", "all_spend": "All AI spend", "prod_spend": "Production spend",
        "dev_spend": "Development spend", "spend_7d": "Spend, last 7 days", "per_1000": "Production spend per 1,000 cases",
        "progress": "Progress and results", "cases": "Cases", "with_result": "With a result", "auto": "Auto-passed",
        "person": "Needs a person", "run_h": "Running time (h)", "untimed": "Batches from before the ledger (no times)",
        "site": "Site", "status": "Status", "share": "Share", "spend": "Spend (USD)", "per1000": "Per 1,000 cases (USD)",
        "wait_h": "Waiting (h)", "elapsed_d": "Elapsed (days)", "started": "Started", "last": "Last progress",
        "date": "Date", "batch": "Batch", "type": "Type", "review_date": "Review date", "reviewed": "Reviewed",
        "correct": "Correct", "error": "Error", "unsure": "Unsure", "two": "Two reviewers", "agree": "Two reviewers agree",
        "disagree": "Two reviewers disagree", "auto_checked": "Auto-pass checked", "auto_wrong": "Auto-pass judged error",
        "note": "Notes", "results_in": "results overwritten by", "total": "Total", "delivered": "Delivered",
        "dev_kind": "Spend by purpose", "dev_module": "Spend by module", "changes": "Changes", "module": "Module",
        "about": "What changed", "took_h": "Took (h)", "released_14": "Changes released, last 14 days", "in_progress": "Changes in progress",
        "median_h": "Start to release, median (h)", "day": "Day", "prod": "Production", "dev": "Development",
        "finished": "Cases finished", "purpose": "Purpose", "hidden": "hidden",
        "foot": "Production and rework batches only. Spend includes history recorded before the ledger.",
        "foot_nomoney": "Production and rework batches only.",
        "st_running": "Running", "st_stopped": "Stopped, needs a restart", "st_paused": "Paused", "st_waiting": "Not started",
        "st_done": "Complete", "st_results_in": "Results in, wrapping up", "st_delivered": "Delivered", "st_closed": "Complete", "st_open": "Not started",
        "st_failed": "Failed", "st_candidate": "In progress", "st_released": "Released", "st_retired": "Retired", "st_abandoned": "Abandoned",
        "p_process_gone": "process gone, needs a restart", "p_disk_full": "disk full", "p_out_of_memory": "stopped: out of memory", "p_program_error": "stopped by an error",
        "p_budget": "waiting for budget", "p_waiting": "waiting for an earlier part", "p_other": "paused",
        "t_production": "production", "t_rework": "rework",
        "kd_experiment": "Experiments", "kd_eval": "Evaluations", "kd_pilot": "Pilots", "kd_drift": "Regression checks",
        "kd_history": "Before the ledger (experiments, evaluations)", "kd_other": "Other", "unassigned": "Not tied to a module",
        "other_reason": "Other", "client": "Client",
    },
    "zh": {
        "Summary": "汇总", "Work packages": "工作包", "Batches": "批次对比", "Development": "开发", "Daily": "每日",
        "title": "生产报表", "updated": "更新于", "cost": "成本", "all_spend": "全部 AI 花费", "prod_spend": "生产花费",
        "dev_spend": "开发花费", "spend_7d": "近 7 天花费", "per_1000": "生产每千个案件花费",
        "progress": "进度与结果", "cases": "案件", "with_result": "已出结果", "auto": "自动通过",
        "person": "需人工复核", "run_h": "运行时长（小时）", "untimed": "接入账本前的批次（无时长）",
        "site": "地区", "status": "状态", "share": "比例", "spend": "花费（美元）", "per1000": "每千案件（美元）",
        "wait_h": "等待（小时）", "elapsed_d": "历时（天）", "started": "开始", "last": "最近进展",
        "date": "日期", "batch": "批次", "type": "类型", "review_date": "复核日期", "reviewed": "已复核",
        "correct": "对", "error": "错", "unsure": "拿不准", "two": "两人复核", "agree": "两人一致",
        "disagree": "两人不一致", "auto_checked": "自动通过抽查", "auto_wrong": "自动通过里判错",
        "note": "备注", "results_in": "结果已被覆盖：", "total": "合计", "delivered": "已交付",
        "dev_kind": "花费按用途", "dev_module": "花费按模块", "changes": "改动", "module": "模块",
        "about": "改了什么", "took_h": "用时（小时）", "released_14": "近 14 天上线的改动", "in_progress": "开发中的改动",
        "median_h": "从开始到上线，中位数（小时）", "day": "日期", "prod": "生产", "dev": "开发",
        "finished": "出结果的案件", "purpose": "用途", "hidden": "已隐藏",
        "foot": "只统计生产和返工批次。花费含接入账本前的历史记录。",
        "foot_nomoney": "只统计生产和返工批次。",
        "st_running": "运行中", "st_stopped": "已停下，需要重启", "st_paused": "暂停", "st_waiting": "未开始",
        "st_done": "已完成", "st_results_in": "结果已出，收尾中", "st_delivered": "已交付", "st_closed": "已完成", "st_open": "未开始",
        "st_failed": "失败", "st_candidate": "开发中", "st_released": "已上线", "st_retired": "已下线", "st_abandoned": "已放弃",
        "p_process_gone": "进程已不在，需要重启", "p_disk_full": "磁盘空间不足", "p_out_of_memory": "内存不足被停下", "p_program_error": "程序出错停下",
        "p_budget": "等待预算批准", "p_waiting": "等前一部分跑完", "p_other": "已暂停",
        "t_production": "生产", "t_rework": "返工",
        "kd_experiment": "实验", "kd_eval": "评测", "kd_pilot": "试跑", "kd_drift": "回归检查",
        "kd_history": "接入账本前（实验、评测）", "kd_other": "其他", "unassigned": "未归到模块",
        "other_reason": "其他", "client": "客户",
    },
}


class _Text:
    def __init__(self, view: dict, lang: str, hide: set[str]):
        self.v, self.lang, self.hide = view, lang, hide
        self.w = WORDS[lang]
        self.clients: dict[str, str] = {}

    def __call__(self, k: str) -> str:
        return self.w.get(k, k)

    def label(self, project: str, key: str) -> str:
        """A project word ([labels] in its reins.toml) in this language; never CJK in English."""
        l = (self.v["labels"].get(project) or {}).get(key)
        s = (l.get(self.lang) or l.get("en") or l.get("zh")) if isinstance(l, dict) else key
        return self.clean(s or key)

    def clean(self, s: str, fallback: str = "") -> str:
        if self.lang == "en" and s and CJK.search(s):
            return fallback or CJK.sub("", s).strip(" ·-") or "?"
        return s

    def name(self, text: str, fallback: str, sites: list[str]) -> str:
        """A batch name; with --hide site, every site word in it (also inside a batch id) becomes its client label."""
        out = self.clean(text, fallback)
        if "site" in self.hide:
            for part in sorted(set(sites), key=len, reverse=True):
                if part:
                    out = re.sub(re.escape(part), self.site(part), out, flags=re.I)
        return out

    def site(self, part: str) -> str:
        if "site" not in self.hide:
            return _title(part)
        if part not in self.clients:
            self.clients[part] = f"{self('client')} {chr(65 + len(self.clients))}"
        return self.clients[part]


def _title(s: str) -> str:
    return " ".join(w.upper() if re.fullmatch(r"wp\d+(wp\d+)*", w) else w[:1].upper() + w[1:]
                    for w in re.sub(r"(wp\d+)(?=wp)", r"\1 ", s or "").replace("-", " ").split())


def _scope_parts(scope: str, n: int) -> list[str]:
    parts = scope.split("-", n - 1) if n > 1 else [scope]
    return parts + [""] * (n - len(parts))


def build(out: Path, lang: str = "en", hide: set[str] | None = None, con=None) -> dict:
    """Write the workbook; returns what it wrote. The numbers are production.view's."""
    from openpyxl import Workbook
    from openpyxl.formatting.rule import DataBarRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from . import production
    if lang not in WORDS:
        raise ReinsError(f"--lang {lang}: one of {sorted(WORDS)}")
    hide = set(hide or ())
    if hide - {"site", "spend"}:
        raise ReinsError(f"--hide {sorted(hide - {'site', 'spend'})}: only site, spend")
    if con is None:
        with closing(connect()) as c:
            v = production.view(c)
    else:
        v = production.view(con)
    T = _Text(v, lang, hide)
    money = "spend" not in hide
    wb = Workbook()
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="3F3D38")
    sect_font = Font(bold=True, size=12, color="9C4527")
    thin = Border(bottom=Side(style="thin", color="D6D3C9"))
    total_font = Font(bold=True)
    FMT = {"n": "#,##0", "pct": "0.0%", "usd": "$#,##0.00", "h": "#,##0.0", "d": "#,##0.0", "date": "yyyy-mm-dd", "text": "@"}

    def table(ws, cols: list[tuple[str, str]], rows: list[list], start_row: int = 1, total: list | None = None, bars: tuple = ()):
        """cols: (header, format). Writes a header, rows, an optional total row; filter, freeze, widths."""
        for j, (h, _) in enumerate(cols, 1):
            c = ws.cell(row=start_row, column=j, value=h)
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for i, r in enumerate(rows + ([total] if total else []), start_row + 1):
            for j, ((_, f), val) in enumerate(zip(cols, r), 1):
                c = ws.cell(row=i, column=j, value=val)
                c.number_format = FMT.get(f, "General")
                c.border = thin
                if total and i == start_row + len(rows) + 1:
                    c.font = total_font
        last = start_row + len(rows) + (1 if total else 0)
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(cols))}{start_row + len(rows)}"
        ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
        for j, (h, f) in enumerate(cols, 1):
            longest = max([len(str(h)) * (1.6 if lang == "zh" else 1)] +
                          [len(str(r[j - 1])) for r in rows if j - 1 < len(r) and r[j - 1] is not None and f == "text"] + [8])
            ws.column_dimensions[get_column_letter(j)].width = min(46, longest + 2)
        ws.row_dimensions[start_row].height = 32
        for j in bars:
            col = get_column_letter(j)
            ws.conditional_formatting.add(f"{col}{start_row + 1}:{col}{start_row + len(rows)}",
                                          DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="5FBF8F"))
        return last

    tot, dev = v["totals"], v["dev"]
    # --- Summary
    ws = wb.active
    ws.title = T("Summary")
    ws["A1"] = T("title"); ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = f"{T('updated')} {v['at'].replace('T', ' ')}"; ws["A2"].font = Font(color="6B6963")
    r = 4
    blocks = []
    if money:
        blocks.append((T("cost"), [(T("all_spend"), round(tot["spent"] + dev["spent"], 2), "usd"), (T("prod_spend"), tot["spent"], "usd"),
                                   (T("dev_spend"), dev["spent"], "usd"), (T("spend_7d"), round(tot["spent_7d"] + dev["spent_7d"], 2), "usd"),
                                   (T("per_1000"), (tot["per_case"] or 0) * 1000, "usd")]))
    got = tot["with_result"] or tot["finished"]
    blocks.append((T("progress"), [(T("Work packages"), tot["packages"], "n"), (T("cases"), tot["n_cases"], "n"),
                                    (T("with_result"), got, "n"), (f"{T('with_result')} %", got / tot["n_cases"] if tot["n_cases"] else None, "pct"),
                                    (T("auto"), tot["good"], "n"), (f"{T('auto')} %", tot["good"] / tot["with_result"] if tot["with_result"] else None, "pct"),
                                    (T("person"), tot["with_result"] - tot["good"], "n"), (T("run_h"), tot["run_h"], "h"),
                                    (T("untimed"), tot["unknown_parts"], "n")]))
    for title, items in blocks:
        ws.cell(row=r, column=1, value=title).font = sect_font
        r += 1
        for k, val, f in items:
            ws.cell(row=r, column=1, value=k)
            c = ws.cell(row=r, column=2, value=val)
            c.number_format, c.font = FMT[f], Font(bold=True)
            r += 1
        r += 1
    ws.cell(row=r, column=1, value=T("foot" if money else "foot_nomoney")).font = Font(italic=True, color="9A978E")
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 18

    # --- Work packages
    ws = wb.create_sheet(T("Work packages"))
    scope_names = (v["batches"][0]["scope_names"] if v["batches"] else ["scope"])
    cols = [(T("site"), "text")] + [(T.label(v["packages"][0]["project"], n) if v["packages"] else n, "text") for n in scope_names[1:]] + \
           [(T("status"), "text"), (T("cases"), "n"), (T("with_result"), "n"), (f"{T('with_result')} %", "pct"),
            (T("auto"), "n"), (f"{T('auto')} %", "pct"), (T("person"), "n")]
    if money:
        cols += [(T("spend"), "usd"), (T("per1000"), "usd")]
    cols += [(T("run_h"), "h"), (T("wait_h"), "h"), (T("elapsed_d"), "d"), (T("started"), "text"), (T("last"), "text")]
    rows = []
    for p in v["packages"]:
        parts = _scope_parts(p["scope"], len(scope_names))
        got = p["with_result"] or p["finished"]
        status = T("st_" + p["status"]) + (f" ({T('p_' + p['pause'])})" if p.get("pause") and p["pause"] != "other" else "")
        row = [T.site(parts[0])] + [_title(x) for x in parts[1:]] + [status, p["n_cases"], got, got / p["n_cases"] if p["n_cases"] else None,
                                                                    p["good"], p["good"] / p["with_result"] if p["with_result"] else None,
                                                                    p["with_result"] - p["good"]]
        if money:
            row += [p["spent"], (p["per_case"] or 0) * 1000 if p["per_case"] is not None else None]
        row += [p["run_h"] if not (p["unknown_parts"] and not p["run_h"]) else None, p["wait_h"], round(p["calendar_h"] / 24, 1),
                (p["start"] or "")[:10], (p["last_activity"] or "")[:16].replace("T", " ")]
        rows.append(row)
    n_fixed = len(scope_names) + 1
    tr = [T("total")] + [""] * (n_fixed - 1) + [tot["n_cases"], tot["with_result"], tot["with_result"] / tot["n_cases"] if tot["n_cases"] else None,
                                                tot["good"], tot["good"] / tot["with_result"] if tot["with_result"] else None, tot["with_result"] - tot["good"]]
    if money:
        tr += [tot["spent"], (tot["per_case"] or 0) * 1000]
    tr += [tot["run_h"], None, None, "", ""]
    table(ws, cols, rows, total=tr, bars=(n_fixed + 3, n_fixed + 5))

    # --- Batches
    ws = wb.create_sheet(T("Batches"))
    for g in v["batches"]:
        P = g["project"]
        sn = g["scope_names"]
        sites = [b["scope"].get(sn[0], "") for b in g["rows"]]
        cols = [(T("date"), "text"), (T("site"), "text")] + [(T.label(P, n), "text") for n in sn[1:]] + \
               [(T("batch"), "text"), (T("type"), "text"), (T("status"), "text"), (T("cases"), "n"), (T("with_result"), "n")]
        dord = g.get("delivered_order") or []
        if dord:                                      # what each case was delivered as, first: the comparison that matters
            cols += [(T("delivered"), "n")] + [(T.label(P, x), "n") for x in dord] + [(f"{T.label(P, x)} %", "pct") for x in dord]
        cols += [(T.label(P, x), "n") for x in g["values"]] + [(f"{T.label(P, x)} %", "pct") for x in g["values"]]
        cols += [(T.label(P, x), "n") for x in g["details"]] + [(f"{T.label(P, x)} %", "pct") for x in g["details"]]
        cols += [(T("review_date"), "text"), (T("reviewed"), "n"), (T("correct"), "n"), (T("error"), "n"), (T("unsure"), "n"),
                 (f"{T('correct')} %", "pct"), (f"{T('error')} %", "pct"), (f"{T('unsure')} %", "pct"),
                 (T("agree"), "n"), (T("disagree"), "n"), (T("auto_checked"), "n"), (T("auto_wrong"), "n")]
        if money:
            cols += [(T("spend"), "usd"), (T("per1000"), "usd")]
        cols += [(T("run_h"), "h"), (T("note"), "text")]
        rows = []
        for b in sorted(g["rows"], key=lambda b: b["date"], reverse=True):             # newest first
            parts = [b["scope"].get(k, "") for k in sn]
            res = b.get("with_result")
            pct = lambda n, d: (n / d) if (n is not None and d) else None
            row = [b["date"], T.site(parts[0])] + [_title(x) for x in parts[1:]] + \
                  [T.name(b["name"], b["batch_id"], sites), T("t_" + b["type"]),
                   T("st_" + b.get("state", b["status"])) + (f" ({T('p_' + b['pause'])})" if b.get("pause") and b["pause"] != "other" else ""), b["n"], res]
            if dord:
                dv, dn = b.get("delivered") or {}, b.get("delivered_n")
                row += [dn] + [dv.get(x, 0) if dn is not None else None for x in dord] + \
                       [pct(dv.get(x, 0), dn) if dn else None for x in dord]
            vals, dets = b.get("values"), b.get("details")
            row += [(vals or {}).get(x, 0) if vals else None for x in g["values"]] + [pct((vals or {}).get(x, 0), res) if vals else None for x in g["values"]]
            row += [(dets or {}).get(x, 0) if vals else None for x in g["details"]] + [pct((dets or {}).get(x, 0), res) if vals else None for x in g["details"]]
            rv = b.get("review")
            row += ([rv["date"], rv["n"], rv["correct"], rv["wrong"], rv["unsure"], pct(rv["correct"], rv["n"]), pct(rv["wrong"], rv["n"]),
                     pct(rv["unsure"], rv["n"]), rv["agree"], rv["double"] - rv["agree"], rv["good_reviewed"], rv["good_wrong"]]
                    if rv else [None] * 12)
            if money:
                row += [b["spent"], b["per_case"] * 1000 if b["per_case"] is not None else None]
            note = f"{T('results_in')} {T.name(b['results_in'], '?', sites)}" if b.get("results_in") else ""
            row += [b["run_h"], note]
            rows.append(row)
        table(ws, cols, rows)
        break                                    # one project per workbook: the first (the board shows the others)

    # --- Development
    ws = wb.create_sheet(T("Development"))
    r = 1
    items = [(T("released_14"), dev["released_14d"], "n"), (T("in_progress"), dev["in_progress"], "n"), (T("median_h"), dev["median_h"], "h")]
    if money:
        items = [(T("dev_spend"), dev["spent"], "usd"), (T("spend_7d"), dev["spent_7d"], "usd")] + items
    for k, val, f in items:
        ws.cell(row=r, column=1, value=k)
        c = ws.cell(row=r, column=2, value=val)
        c.number_format, c.font = FMT[f], Font(bold=True)
        r += 1
    proj = v["packages"][0]["project"] if v["packages"] else ""
    if money:
        r += 1
        ws.cell(row=r, column=1, value=T("dev_kind")).font = sect_font
        r += 1
        for x in dev["by_kind"]:
            ws.cell(row=r, column=1, value=T("kd_" + x["kind"])); ws.cell(row=r, column=2, value=x["usd"]).number_format = FMT["usd"]; r += 1
        r += 1
        ws.cell(row=r, column=1, value=T("dev_module")).font = sect_font
        r += 1
        for x in dev["by_module"]:
            ws.cell(row=r, column=1, value=T("unassigned") if x["module"] == "unassigned" else T.label(proj, x["module"]))
            ws.cell(row=r, column=2, value=x["usd"]).number_format = FMT["usd"]; r += 1
    r += 1
    ws.cell(row=r, column=1, value=T("changes")).font = sect_font
    cols = [(T("module"), "text"), (T("status"), "text"), (T("started"), "text"), (T("took_h"), "h")]
    if lang != "en":
        cols.insert(1, (T("about"), "text"))
    if money:
        cols.append((T("spend"), "usd"))
    rows = []
    for c in dev["changes"]:
        row = [T.label(c["project"], c["module"]), T("st_" + c["status"]), (c["start"] or "")[:10], c["hours"]]
        if lang != "en":
            row.insert(1, T.name(c["about"], "", [k for k in T.clients] + [b["scope"].get(sn[0], "") for g in v["batches"] for b in g["rows"]]))
        if money:
            row.append(c["spent"])
        rows.append(row)
    table(ws, cols, rows, start_row=r + 1)
    ws.column_dimensions["A"].width = 40
    ws.freeze_panes = None

    # --- Daily
    ws = wb.create_sheet(T("Daily"))
    devd = {d["day"]: d["usd"] for d in dev["daily"]}
    cols = [(T("day"), "text"), (T("finished"), "n")] + ([(f"{T('spend')} · {T('prod')}", "usd"), (f"{T('spend')} · {T('dev')}", "usd")] if money else [])
    rows = [[d["day"], d["n"]] + ([d["usd"], devd.get(d["day"], 0)] if money else []) for d in v["daily"]]
    table(ws, cols, rows)

    if lang == "en":                                                      # the promise: no CJK anywhere in the English file
        for sh in wb.worksheets:
            for row in sh.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and CJK.search(c.value):
                        raise ReinsError(f"English report would contain CJK text in {sh.title}!{c.coordinate}: {c.value[:40]!r}")
    if "site" in hide:                                                    # the promise: no site name anywhere
        names = {p["scope"].split("-")[0] for p in v["packages"]}
        pat = re.compile("|".join(re.escape(n) for n in names if len(n) > 2), re.I) if names else None
        for sh in wb.worksheets:
            for row in sh.iter_rows():
                for c in row:
                    if pat and isinstance(c.value, str) and pat.search(c.value):
                        raise ReinsError(f"report with hidden sites would name one in {sh.title}!{c.coordinate}: {c.value[:40]!r}")
    out.parent.mkdir(parents=True, exist_ok=True)
    tmp = out.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    tmp.replace(out)
    return {"path": str(out), "sheets": wb.sheetnames, "packages": len(v["packages"]),
            "batches": sum(len(g["rows"]) for g in v["batches"][:1]), "lang": lang, "hide": sorted(hide)}


def settings() -> dict:
    return {"remote": "", "lang": "en", "hide": [], "out": "", **(config().get("report") or {})}


def default_out(lang: str) -> Path:
    s = settings()
    return Path(s["out"]) if s.get("out") else home() / "reports" / f"production-{lang}.xlsx"


def publish(lang: str | None = None, hide: set[str] | None = None) -> dict:
    """Build, then copy to [report] remote with rclone (the remote file is replaced; its share link stays)."""
    s = settings()
    lang = lang or s["lang"]
    hide = set(hide if hide is not None else s["hide"])
    out = default_out(lang)
    r = build(out, lang, hide)
    if not s["remote"]:
        raise ReinsError(f"wrote {out}; no [report] remote in {home() / 'config.toml'} to copy it to")
    p = subprocess.run(["rclone", "copyto", str(out), s["remote"]], capture_output=True, text=True, timeout=300)
    if p.returncode:
        raise ReinsError(f"rclone copyto {s['remote']} failed: {p.stderr.strip()[-400:]}")
    return {**r, "remote": s["remote"]}


# ---------- the production page on Cloudflare Workers (deploy/cloudflare) ----------
#
#   reins report deploy        create the KV store, set the access codes, deploy the Worker (once, or after a page change)
#   reins report push          send the newest data and workbook (reins-report-push.timer, every 5 minutes)
#   reins report link [--new]  the viewer link; --new adds another access code (one per person, so one can be revoked)
#   reins report revoke CODE   drop an access code
#
# $REINS_HOME/config.toml [report.cloudflare]: name = "reins-production" (the Worker), lang = "en", hide = [],
# wrangler = "npx --yes wrangler@4", login_home = "" (a directory: wrangler logs in there, `reins report login`). State (KV id, URL, codes) lives in $REINS_HOME/reports/cloudflare/state.json (0600).

import json as _json
import os as _os
import secrets as _secrets
import shlex as _shlex
import shutil as _shutil

TEMPLATE = Path(__file__).resolve().parents[1] / "deploy" / "cloudflare"
UI = Path(__file__).with_name("board_ui")


def _cf() -> dict:
    return {"name": "reins-production", "lang": "en", "hide": [], "wrangler": "npx --yes wrangler@4", "login_home": "",
            **(settings().get("cloudflare") or {})}


def _env() -> dict:
    """wrangler's environment: with login_home set, its own login there (HOME), so this machine's other wrangler
    projects keep theirs; the npm cache stays the user's."""
    env = dict(_os.environ)
    h = _cf()["login_home"]
    if h:
        Path(h).mkdir(parents=True, exist_ok=True)
        env["NPM_CONFIG_CACHE"] = env.get("NPM_CONFIG_CACHE") or str(Path.home() / ".npm")
        env["HOME"] = str(Path(h).expanduser())
    return env


def _dir() -> Path:
    d = home() / "reports" / "cloudflare"
    d.mkdir(parents=True, exist_ok=True)
    return d


def _state() -> dict:
    f = _dir() / "state.json"
    return _json.loads(f.read_text()) if f.is_file() else {}


def _save(st: dict) -> None:
    f = _dir() / "state.json"
    f.write_text(_json.dumps(st, indent=2))
    _os.chmod(f, 0o600)


def _wrangler(*args, stdin: str | None = None, cwd: Path | None = None) -> str:
    cmd = _shlex.split(_cf()["wrangler"]) + list(args)
    p = subprocess.run(cmd, input=stdin, capture_output=True, text=True, timeout=600, cwd=cwd or _dir(), env=_env())
    if p.returncode:
        raise ReinsError(f"wrangler {' '.join(args[:3])} failed: {(p.stderr or p.stdout).strip()[-600:]}")
    return p.stdout + p.stderr


def _wrangler_out(*args) -> str:
    """stdout only (JSON answers), not the banner wrangler writes to stderr."""
    cmd = _shlex.split(_cf()["wrangler"]) + list(args)
    p = subprocess.run(cmd, capture_output=True, text=True, timeout=600, cwd=_dir(), env=_env())
    if p.returncode:
        raise ReinsError(f"wrangler {' '.join(args[:3])} failed: {(p.stderr or p.stdout).strip()[-600:]}")
    out = p.stdout
    return out[out.index("["):] if "[" in out else out


def _hide_sites(text: str, names: list[str]) -> str:
    """Every site word (also inside ids and names) -> Client A, B, ... in order of first appearance."""
    for i, n in enumerate(sorted({x for x in names if len(x) > 2}, key=lambda x: names.index(x))):
        text = re.sub(re.escape(n), f"Client {chr(65 + i)}", text, flags=re.I)
    return text


def deploy() -> dict:
    cf, st = _cf(), _state()
    d = _dir()
    if not st.get("kv_id"):
        title = f"{cf['name']}-data"
        try:
            _wrangler("kv", "namespace", "create", title)
        except ReinsError as e:
            if "already exists" not in str(e):
                raise
        found = [n["id"] for n in _json.loads(_wrangler_out("kv", "namespace", "list")) if n.get("title") == title]
        if not found:
            raise ReinsError(f"KV namespace {title} not found after creating it")
        st["kv_id"] = found[0]
    pub = d / "public"
    if pub.exists():
        _shutil.rmtree(pub)
    pub.mkdir()
    for f in ("production.js", "production.css", "app.css"):
        _shutil.copy(UI / f, pub / f)
    html = (UI / "production.html").read_text(encoding="utf-8")
    html = html.replace('<html lang="zh">', f'<html lang="{cf["lang"]}" data-lang="{cf["lang"]}" data-xlsx="production.xlsx">')
    html = html.replace("<head>", '<head>\n<meta name="robots" content="noindex">', 1)
    (pub / "production.html").write_text(html, encoding="utf-8")
    _shutil.copy(TEMPLATE / "worker.js", d / "worker.js")
    (d / "wrangler.toml").write_text(
        f'name = "{cf["name"]}"\nmain = "worker.js"\ncompatibility_date = "2026-10-01"\nworkers_dev = true\n\n'
        f'[assets]\ndirectory = "./public"\nbinding = "ASSETS"\nrun_worker_first = true\n\n'
        f'[[kv_namespaces]]\nbinding = "REPORT"\nid = "{st["kv_id"]}"\n')
    out = _wrangler("deploy")
    m = re.search(r"https://[\w.-]+\.workers\.dev", out)
    if m:
        st["url"] = m.group(0)
    if not st.get("codes"):
        st["codes"] = [_secrets.token_urlsafe(24)]
    _save(st)
    _wrangler("secret", "put", "VIEW_TOKENS", stdin=",".join(st["codes"]))
    r = push()
    return {**r, "url": st.get("url"), "kv_id": st["kv_id"]}


def push() -> dict:
    """The newest data (the board's /api/production) and the workbook, into the Worker's KV store."""
    from . import production
    cf, st = _cf(), _state()
    if not st.get("kv_id"):
        raise ReinsError("not deployed yet: reins report deploy")
    hide = set(cf["hide"])
    with closing(connect()) as con:
        v = production.view(con)
        text = _json.dumps(v, ensure_ascii=False, default=str)
        if "site" in hide:
            names = []
            for p in v["packages"]:
                n = p["scope"].split("-")[0]
                if n not in names:
                    names.append(n)
            text = _hide_sites(text, names)
        d = _dir()
        (d / "production.json").write_text(text, encoding="utf-8")
        build(d / "production.xlsx", cf["lang"], hide, con=con)
    _wrangler("kv", "key", "put", "production", "--path", str(d / "production.json"), "--namespace-id", st["kv_id"], "--remote")
    _wrangler("kv", "key", "put", "xlsx", "--path", str(d / "production.xlsx"), "--namespace-id", st["kv_id"], "--remote")
    st["pushed"] = v["at"]
    _save(st)
    return {"path": str(d / "production.xlsx"), "sheets": [], "packages": len(v["packages"]), "batches": 0,
            "remote": st.get("url") or "cloudflare"}


def link(new: bool = False) -> list[str]:
    st = _state()
    if not st.get("url"):
        raise ReinsError("not deployed yet: reins report deploy")
    if new:
        st["codes"] = st.get("codes", []) + [_secrets.token_urlsafe(24)]
        _save(st)
        _wrangler("secret", "put", "VIEW_TOKENS", stdin=",".join(st["codes"]))
    return [f"{st['url']}/?k={c}" for c in st["codes"]]


def revoke(code: str) -> int:
    st = _state()
    left = [c for c in st.get("codes", []) if c != code]
    if len(left) == len(st.get("codes", [])):
        raise ReinsError("no such access code")
    st["codes"] = left
    _save(st)
    _wrangler("secret", "put", "VIEW_TOKENS", stdin=",".join(left) or "-")
    return len(left)


def login() -> None:
    """wrangler login into login_home (opens the browser; the person approves there)."""
    cmd = _shlex.split(_cf()["wrangler"]) + ["login"]
    p = subprocess.run(cmd, cwd=_dir(), env=_env(), timeout=900)
    if p.returncode:
        raise ReinsError("wrangler login did not finish")
