"""The production page (board /production): what a manager looks at. Production and rework batches only, one block per
work package (batch scope): how many cases have a result, what the results are, why the others could not be processed,
what it cost, how long it took (running vs waiting) and when it should be done. No commands, sessions, versions or
paths. Cost = the gateway ledger plus history imported from before the gateway (reins spend import) whose label is the
batch id or the end of its work dir; time = the batch's processes (or stages run under reins), parallel parts counted
once. Batches adopted from before reins have no times: they are counted and shown as such.

A project says where a case's result lives in its reins.toml (optional; without it the page counts finished stages):

    [results]
    file = "out/results.csv"          # relative to the batch work_dir
    key = "case"                      # the case column
    key_pattern = "([0-9]+)$"         # optional: the case id inside the key (SHF_100047 -> 100047)
    value = "lane"                    # the result column
    detail = "category"               # optional: a breakdown of the value
    good = ["auto pass"]              # the values that need nobody
    [[results.reasons]]               # a skip / fail reason -> a word (first match wins)
    match = "no scans|no source"
    zh = "..."                        # any language keys the page offers
    en = "No scans"
    [labels]                          # stage names, result values, details -> words per language
    lane_a = { zh = "...", en = "..." }

A case's result is taken from the newest batch of the work package that has a row for it; rows of cases outside the
batch are ignored (a results file may hold more than its batch: conservation, C1).
"""
from __future__ import annotations

import csv
import datetime as dt
import glob
import json
import os
import re
from pathlib import Path

from . import projects

TYPES = ("production", "rework")
_FILES: dict[str, tuple[float, dict]] = {}


def _cfg(project: str) -> dict:
    for _, c in projects.all_():
        if c.get("project") == project:
            return c
    return {}


def _results(path: Path, spec: dict) -> dict[str, tuple[str, str]]:
    """case_id -> (value, detail), cached by mtime."""
    try:
        mt = path.stat().st_mtime
    except OSError:
        return {}
    key = str(path)
    if key in _FILES and _FILES[key][0] == mt:
        return _FILES[key][1]
    pat = re.compile(spec["key_pattern"]) if spec.get("key_pattern") else None
    out = {}
    try:
        with open(path, newline="", encoding="utf-8", errors="replace") as f:
            for row in csv.DictReader(f):
                k = (row.get(spec.get("key", "case_id")) or "").strip()
                if pat:
                    m = pat.search(k)
                    k = m.group(1) if m else k
                v = (row.get(spec.get("value", "result")) or "").strip()
                if k and v:
                    out[k] = (v, (row.get(spec["detail"]) or "").strip() if spec.get("detail") else "")
    except OSError:
        return {}
    _FILES[key] = (mt, out)
    return out


def _reason_word(reason: str, rules: list[dict]) -> str:
    for i, r in enumerate(rules):
        try:
            if re.search(r["match"], reason or "", re.I):
                return f"r{i}"
        except re.error:
            continue
    return "other"


def _pause_word(reason: str | None) -> str:
    """reins' own pause reasons, in a manager's words (the words are keys the page translates)."""
    r = (reason or "").lower()
    if "disk" in r:
        return "disk_full"
    if "memory" in r or "killed" in r:
        return "out_of_memory"
    if "exited" in r or "error" in r:
        return "program_error"
    if "cap" in r or "spend" in r or "budget" in r:
        return "budget"
    if r.startswith("等") or "wait" in r:
        return "waiting"
    return "other"


def _part_state(con, b) -> tuple[str, str | None]:
    """(state, the newest sign of work: an outcome or a process start/end)."""
    from .board import _activity, ACTIVE_H
    if b["status"] in ("done", "closed"):
        return "done", None
    if b["status"] == "paused":
        return "paused", None
    if b["status"] == "open":
        return "waiting", None
    a = _activity(con, b["batch_id"])
    recent = a["progress"] and (dt.datetime.now() - dt.datetime.fromisoformat(a["progress"])).total_seconds() < ACTIVE_H * 3600
    return ("running" if a["alive"] or recent else "stopped"), a["at"]


def _abandoned(con, bid: str) -> bool:
    return bool(con.execute("SELECT 1 FROM batch_event WHERE batch_id=? AND event='abandoned'", (bid,)).fetchone())


def _iso(d: dt.datetime) -> str:
    return d.isoformat(timespec="seconds")


def _union_h(intervals) -> float:
    """Hours covered by at least one interval (parallel parts are not counted twice)."""
    pts = sorted((dt.datetime.fromisoformat(a), dt.datetime.fromisoformat(b)) for a, b in intervals if a and b and b > a)
    total, cur = 0.0, None
    for a, b in pts:
        if cur and a <= cur[1]:
            cur[1] = max(cur[1], b)
            continue
        if cur:
            total += (cur[1] - cur[0]).total_seconds()
        cur = [a, b]
    if cur:
        total += (cur[1] - cur[0]).total_seconds()
    return total / 3600


def _intersect(a: list[tuple[str, str]], b: list[tuple[str, str]]) -> list[tuple[str, str]]:
    """The parts of intervals a that fall inside intervals b (a stage's span while something of its batch ran)."""
    out = []
    for s1, e1 in a:
        for s2, e2 in b:
            if not (s1 and e1 and s2 and e2):
                continue
            lo, hi = max(s1, s2), min(e1, e2)
            if hi > lo:
                out.append((lo, hi))
    return out


def _run_intervals(con, bid: str, now: str) -> list[tuple[str, str]]:
    """When the batch had something running: its processes (start to end; a live one to now; a lost one to its last
    outcome), else its stages run under reins. Stages adopted from before reins have no times: empty."""
    from . import runner
    out = []
    for p in con.execute("SELECT * FROM process WHERE batch_id=? ORDER BY id", (bid,)).fetchall():
        end = p["ended"]
        if not end:
            if p["state"] == "running" and (runner._alive(p["pgid"]) or runner._alive(p["supervisor_pid"])):
                end = now
            else:
                end = con.execute("SELECT MAX(at) FROM case_event WHERE batch_id=? AND at>=?", (bid, p["started"])).fetchone()[0]
        out.append((p["started"], end))
    if not out:
        out = [(a, b) for _, a, b in _stage_intervals(con, bid, now)]
    return out


def _stage_intervals(con, bid: str, now: str) -> list[tuple[str, str, str]]:
    out = []
    for r in con.execute("SELECT stage, started, ended, status FROM batch_stage WHERE batch_id=? AND started IS NOT NULL"
                         " AND COALESCE(note,'')<>'adopted'", (bid,)).fetchall():
        end = r["ended"] or con.execute("SELECT MAX(at) FROM case_event WHERE batch_id=? AND stage=?", (bid, r["stage"])).fetchone()[0]
        out.append((r["stage"], r["started"], end or r["started"]))
    return out


def _imported_spend(con, rows) -> dict[str, list[tuple[str, str, float]]]:
    """History imported from before the gateway (reins spend import), by batch: its label is the batch id or the end of
    the batch's work dir (e.g. 'sheffield/wp3-batch2-e2e'). -> {batch_id: [(purpose, day, usd)]}."""
    if not con.execute("SELECT 1 FROM sqlite_master WHERE name='spend_import'").fetchone():
        return {}
    by_label = {}
    for b in rows:
        by_label[b["batch_id"]] = b["batch_id"]
        if b["work_dir"]:
            wd = b["work_dir"].rstrip("/")
            by_label.setdefault(wd, b["batch_id"])
            parts = wd.split("/")
            for i in range(1, len(parts)):
                by_label.setdefault("/".join(parts[i:]), b["batch_id"])
    out: dict[str, list] = {}
    for label, purpose, day, usd in con.execute("SELECT batch_label, purpose, day, SUM(amount) FROM spend_import"
                                                " WHERE batch_label IS NOT NULL GROUP BY 1,2,3"):
        bid = by_label.get(label.strip().rstrip("/"))
        if bid:
            out.setdefault(bid, []).append((purpose, day, usd or 0))
    return out


def view(con) -> dict:
    from .board import _eta
    all_prod = con.execute("SELECT * FROM batch WHERE type IN (%s) ORDER BY created" % ",".join("?" * len(TYPES)), TYPES).fetchall()
    rows = [b for b in all_prod if b["status"] != "failed" and not _abandoned(con, b["batch_id"])]   # failed = superseded or never ran
    dropped = [b for b in all_prod if b not in rows]
    ids_all = {b["batch_id"] for b in all_prod}
    ledger: dict[str, list] = {}                                                   # batch -> [(stage, day, usd)]
    for bid, stage, day, usd in con.execute("SELECT batch_id, stage, substr(at,1,10), SUM(amount) FROM spend GROUP BY 1,2,3"):
        if bid in ids_all:
            ledger.setdefault(bid, []).append((stage, day, usd or 0))
    first_spend = {r[0]: r[1] for r in con.execute("SELECT batch_id, MIN(at) FROM spend GROUP BY batch_id")}
    imported = _imported_spend(con, all_prod)
    ai_all = (con.execute("SELECT COALESCE(SUM(amount),0) FROM spend").fetchone()[0]
              + (con.execute("SELECT COALESCE(SUM(amount),0) FROM spend_import").fetchone()[0]
                 if con.execute("SELECT 1 FROM sqlite_master WHERE name='spend_import'").fetchone() else 0))
    spend_daily: dict[str, float] = {}
    labels, reasons_cfg, packages = {}, {}, {}
    for b in rows:
        packages.setdefault((b["project"], b["scope"]), []).append(b)
    now = dt.datetime.now()
    day_ago = (now - dt.timedelta(hours=24)).isoformat(timespec="seconds")
    daily: dict[str, int] = {}
    out = []
    for (project, scope), parts in packages.items():
        cfg = _cfg(project)
        spec = cfg.get("results") or {}
        rules = spec.get("reasons") or []
        if project not in labels:
            labels[project] = cfg.get("labels") or {}
            reasons_cfg[project] = [{k: v for k, v in r.items() if k != "match"} for r in rules]
        good_vals = set(spec.get("good") or [])
        ids = [p["batch_id"] for p in parts]
        q = ",".join("?" * len(ids))
        members: dict[str, set] = {p["batch_id"]: set() for p in parts}
        for bid, c in con.execute(f"SELECT batch_id, case_id FROM batch_case WHERE batch_id IN ({q})", ids):
            members[bid].add(c)
        cases = set().union(*members.values())
        stage_order = [r[0] for r in con.execute("SELECT stage FROM batch_stage WHERE batch_id=? ORDER BY ord", (parts[-1]["batch_id"],))]
        # newest outcome per (stage, case) across the parts; the first time a case finished the last stage
        outcome: dict[tuple[str, str], tuple[str, str]] = {}
        last_stage = stage_order[-1] if stage_order else None
        first_final: dict[str, str] = {}
        last_at = None
        for at, stage, case, status, reason in con.execute(
                f"SELECT at, stage, case_id, status, reason FROM case_event WHERE batch_id IN ({q}) AND status<>'started' ORDER BY id", ids):
            outcome[(stage, case)] = (status, reason or "")
            last_at = at if not last_at or at > last_at else last_at
            if stage == last_stage and case not in first_final:
                first_final[case] = at
        result: dict[str, tuple[str, str]] = {}
        if spec.get("file"):
            for p in sorted(parts, key=lambda p: p["created"], reverse=True):
                if not p["work_dir"]:
                    continue
                for c, v in _results(Path(p["work_dir"]) / spec["file"], spec).items():
                    if c in members[p["batch_id"]] and c not in result:
                        result[c] = v
        finished = {c for c in cases if c in result or (last_stage and (last_stage, c) in outcome)}
        by_value: dict[str, dict] = {}
        for c, (v, d) in result.items():
            e = by_value.setdefault(v, {"value": v, "n": 0, "good": v in good_vals, "details": {}})
            e["n"] += 1
            if d:
                e["details"][d] = e["details"].get(d, 0) + 1
        results = sorted(({**e, "details": [{"value": k, "n": n} for k, n in sorted(e["details"].items())]}
                          for e in by_value.values()), key=lambda e: (not e["good"], -e["n"]))
        # cases with no result: did a stage give up on them, and why
        why: dict[str, int] = {}
        for c in cases - set(result):
            bad = next((outcome[(s, c)][1] for s in stage_order if outcome.get((s, c), ("",))[0] in ("skipped", "failed")), None)
            if bad is not None:
                w = _reason_word(bad, rules)
                why[w] = why.get(w, 0) + 1
        stages = [{"stage": s, "done": sum(1 for c in cases if (s, c) in outcome)} for s in stage_order]
        now_iso = _iso(now)
        stage_names = set(stage_order)
        cost_stage: dict[str, float] = {}
        stage_iv: dict[str, list] = {}
        run_iv, unknown_time, first_seen = [], 0, []
        part_rows, etas = [], []
        for p in parts:
            pid = p["batch_id"]
            p_spent = 0.0
            for stage, day, usd in ledger.get(pid, []):
                p_spent += usd
                cost_stage[stage] = cost_stage.get(stage, 0) + usd
                spend_daily[day] = spend_daily.get(day, 0) + usd
            if pid in first_spend:
                first_seen.append(first_spend[pid])
            for purpose, day, usd in imported.get(pid, []):
                p_spent += usd
                st = (purpose or "").split(":")[0].strip()
                st = st if st in stage_names else "other"
                cost_stage[st] = cost_stage.get(st, 0) + usd
                if day:
                    spend_daily[day] = spend_daily.get(day, 0) + usd
                    first_seen.append(day + "T00:00:00")
            iv = _run_intervals(con, pid, now_iso)
            if not iv:
                unknown_time += 1
            run_iv += iv
            first_seen += [a for a, _ in iv]
            for st, a, b in _stage_intervals(con, pid, now_iso):         # paused or waiting inside a stage is not its time
                stage_iv.setdefault(st, []).extend(_intersect([(a, b)], iv) if iv else [])
            state, seen = _part_state(con, p)
            if seen and state == "running" and (not last_at or seen > last_at):
                last_at = seen
            n = len(members[p["batch_id"]])
            got = sum(1 for c in members[p["batch_id"]] if c in finished)
            eta = None
            if state == "running":
                cur = con.execute("SELECT stage FROM batch_stage WHERE batch_id=? AND status='running' ORDER BY ord LIMIT 1",
                                  (p["batch_id"],)).fetchone()
                if cur:
                    acc = con.execute("SELECT COUNT(DISTINCT case_id) FROM case_event WHERE batch_id=? AND stage=? AND status<>'started'",
                                      (p["batch_id"], cur[0])).fetchone()[0]
                    eta = _eta(con, p["batch_id"], cur[0], n, acc)
                    eta = None if eta == "-" else eta
                    if eta:
                        etas.append(eta)
            dl = con.execute("SELECT version, at, n_rows FROM deliverable WHERE batch_id=? ORDER BY id DESC LIMIT 1", (p["batch_id"],)).fetchone()
            acc = con.execute("SELECT decision, decided_at FROM acceptance WHERE batch_id=? AND decision IS NOT NULL ORDER BY id DESC LIMIT 1",
                              (p["batch_id"],)).fetchone()
            aliases = json.loads(p["aliases"] or "[]")
            part_rows.append({"name": aliases[0] if aliases else p["purpose"][:60], "type": p["type"], "state": state,
                              "pause": _pause_word(p["status_reason"]) if state == "paused" else None,
                              "n": n, "finished": got, "started": p["created"], "ended": p["closed"], "eta": eta,
                              "spent": round(p_spent, 2), "budget": p["spend_cap"] or 0,
                              "run_h": round(_union_h(iv), 1) if iv else None,
                              "delivered": dict(dl) if dl else None, "accepted": dict(acc) if acc else None})
        states = {r["state"] for r in part_rows}
        if states == {"done"}:
            status = "done"
        elif "running" not in states and cases and len(finished) == len(cases):
            status = "results_in"                # every case has a result, but a batch was stopped and never closed
        else:
            status = next(s for s in ("running", "stopped", "paused", "waiting", "done") if s in states)
        pause = next((r["pause"] for r in part_rows if r["state"] == "paused"), None) if status == "paused" else None
        for c, at in first_final.items():
            if c in cases:
                daily[at[:10]] = daily.get(at[:10], 0) + 1
        start = min([p["created"] for p in parts] + first_seen)
        ends = [b for _, b in run_iv if b] + [p["closed"] for p in parts if p["closed"]] + ([last_at] if last_at else [])
        end = now_iso if status in ("running", "stopped", "paused", "waiting") else max(ends + [start])
        cal_h = (dt.datetime.fromisoformat(end) - dt.datetime.fromisoformat(start)).total_seconds() / 3600
        run_h = _union_h(run_iv)
        spent_pkg = sum(r["spent"] for r in part_rows)
        time_cost = {"start": start, "end": end, "calendar_h": round(cal_h, 1), "run_h": round(run_h, 1),
                     "wait_h": round(max(0.0, cal_h - run_h), 1) if not unknown_time else None,
                     "unknown_parts": unknown_time, "budget": round(sum(r["budget"] for r in part_rows), 2),
                     "per_case": round(spent_pkg / len(cases), 4) if cases else None,
                     "per_result": round(spent_pkg / len(result), 4) if result else None,
                     "cost_by_stage": [{"stage": k, "usd": round(v, 2)} for k, v in
                                       sorted(cost_stage.items(), key=lambda kv: (stage_order + ["other"]).index(kv[0])
                                              if kv[0] in stage_order + ["other"] else 99) if v > 0.005],
                     "time_by_stage": [{"stage": k, "h": round(_union_h(stage_iv[k]), 1)} for k in stage_order
                                       if k in stage_iv and _union_h(stage_iv[k]) >= 0.05]}
        retry = 0.0                              # money spent on tries of this package that were superseded
        for d in dropped:
            if (d["project"], d["scope"]) == (project, scope):
                for stage, day, usd in ledger.get(d["batch_id"], []) + [(pp, dd, uu) for pp, dd, uu in imported.get(d["batch_id"], [])]:
                    retry += usd
                    if day:
                        spend_daily[day] = spend_daily.get(day, 0) + usd
        spent_pkg += retry
        time_cost["spent_retry"] = round(retry, 2)
        time_cost["per_case"] = round(spent_pkg / len(cases), 4) if cases else None
        time_cost["per_result"] = round(spent_pkg / len(result), 4) if result else None
        open_issues = con.execute(f"SELECT COUNT(*) FROM issue WHERE batch_id IN ({q}) AND status IN ('open','in_progress','fixed')",
                                  ids).fetchone()[0] if con.execute("SELECT 1 FROM sqlite_master WHERE name='issue'").fetchone() else 0
        out.append({"project": project, "scope": scope, "status": status, "pause": pause, "n_cases": len(cases),
                    "finished": len(finished), "with_result": len(result),
                    "good": sum(e["n"] for e in results if e["good"]), "results": results,
                    "why": sorted(({"word": k, "n": n} for k, n in why.items()), key=lambda x: -x["n"]),
                    "stages": stages, "parts": part_rows, "started": min(p["created"] for p in parts),
                    "last_activity": last_at, "eta": etas[0] if len(etas) == 1 else (max(etas, key=_eta_minutes) if etas else None),
                    "today": sum(1 for c, at in first_final.items() if at > day_ago and c in cases),
                    "spent": round(spent_pkg, 2), "open_issues": open_issues, **time_cost})
    order = {"running": 0, "stopped": 1, "paused": 2, "waiting": 3, "results_in": 4, "done": 5}
    out.sort(key=lambda p: (order[p["status"]], -(dt.datetime.fromisoformat(p["last_activity"] or p["started"]).timestamp())))
    days = [(now.date() - dt.timedelta(days=i)).isoformat() for i in range(13, -1, -1)]
    tot = {k: sum(p[k] for p in out) for k in ("n_cases", "finished", "with_result", "good", "today", "spent")}
    tot["packages"] = len(out)
    tot["spent"] = round(tot["spent"], 2)
    tot["ai_all"] = round(ai_all, 2)
    tot["per_case"] = round(tot["spent"] / tot["n_cases"], 4) if tot["n_cases"] else None
    tot["per_result"] = round(tot["spent"] / tot["with_result"], 4) if tot["with_result"] else None
    tot["run_h"] = round(sum(p["run_h"] for p in out), 1)
    tot["unknown_parts"] = sum(p["unknown_parts"] for p in out)
    week = {(now.date() - dt.timedelta(days=i)).isoformat() for i in range(7)}
    tot["spent_7d"] = round(sum(v for d, v in spend_daily.items() if d in week), 2)
    tot["running"] = sum(1 for p in out if p["status"] == "running")
    dev = development(con, ids_all, imported, days)
    table = batch_table(con, rows, ledger, imported)
    return {"at": now.isoformat(timespec="seconds"), "totals": tot, "dev": dev, "batches": table, "packages": out, "labels": labels,
            "reasons": reasons_cfg, "daily": [{"day": d, "n": daily.get(d, 0), "usd": round(spend_daily.get(d, 0), 2)} for d in days]}


def _eta_minutes(s: str) -> float:
    try:
        v, u = s.split()
        return float(v) * (60 if u == "h" else 1)
    except ValueError:
        return 0


def development(con, prod_ids: set, prod_imported: dict, days: list[str]) -> dict:
    """What development cost: every AI call not made for a production or rework batch (experiments, evaluations,
    pilots, drift checks, imported history from before reins that no production batch claims), by kind, by module and
    by change; and how long changes took from start to release."""
    now = dt.datetime.now()
    claimed = {(p, d, round(u, 6)) for v in prod_imported.values() for p, d, u in v}
    by_kind: dict[str, float] = {}
    by_module: dict[str, float] = {}
    by_version: dict[str, float] = {}
    daily: dict[str, float] = {}
    stage_mv = {}
    for bid, mv in con.execute("SELECT batch_id, MIN(module_version) FROM batch_stage WHERE module_version IS NOT NULL"
                               " GROUP BY batch_id HAVING COUNT(DISTINCT module_version)=1"):
        stage_mv[bid] = mv
    btype = {r[0]: r[1] for r in con.execute("SELECT batch_id, type FROM batch")}
    mod_of = {r[0]: r[1] for r in con.execute("SELECT version, module FROM module_version")}
    for bid, mv, day, usd in con.execute("SELECT batch_id, module_version, substr(at,1,10), SUM(amount) FROM spend GROUP BY 1,2,3"):
        if bid in prod_ids or not usd:
            continue
        kind = btype.get(bid, "other")
        by_kind[kind] = by_kind.get(kind, 0) + usd
        v = mv or stage_mv.get(bid)
        if v:
            by_version[v] = by_version.get(v, 0) + usd
        m = mod_of.get(v, "") if v else ""
        by_module[m or "unassigned"] = by_module.get(m or "unassigned", 0) + usd
        daily[day] = daily.get(day, 0) + usd
    if con.execute("SELECT 1 FROM sqlite_master WHERE name='spend_import'").fetchone():
        for label, purpose, day, usd in con.execute("SELECT batch_label, purpose, day, SUM(amount) FROM spend_import GROUP BY 1,2,3"):
            if (purpose, day, round(usd or 0, 6)) in claimed or not usd:
                continue
            by_kind["history"] = by_kind.get("history", 0) + usd
            by_module["unassigned"] = by_module.get("unassigned", 0) + usd
            if day:
                daily[day] = daily.get(day, 0) + usd
    ev: dict[str, dict] = {}
    for v, event, at in con.execute("SELECT version, event, at FROM module_event WHERE event IN ('created','released','abandoned') ORDER BY id"):
        ev.setdefault(v, {}).setdefault(event, at)
    changes = []
    for r in con.execute("SELECT v.*, m.project FROM module_version v JOIN module m ON m.name=v.module"
                         " WHERE v.suffix<>'main' ORDER BY v.created DESC"):          # <module>-main-*: the baseline registered at adoption
        e = ev.get(r["version"], {})
        start = e.get("created") or r["created"]
        end = e.get("released") or e.get("abandoned")
        h = ((dt.datetime.fromisoformat(end) if end else now) - dt.datetime.fromisoformat(start)).total_seconds() / 3600
        changes.append({"project": r["project"], "module": r["module"], "about": r["about"], "status": r["status"],
                        "start": start, "end": end, "hours": round(h, 1), "spent": round(by_version.get(r["version"], 0), 2)})
    recent = (now - dt.timedelta(days=14)).isoformat(timespec="seconds")
    released_14 = [c for c in changes if c["status"] in ("released", "retired") and c["end"] and c["end"] >= recent]
    rel_all = [c for c in changes if c["status"] in ("released", "retired") and c["end"]]
    total = sum(by_kind.values())
    week = {(now.date() - dt.timedelta(days=i)).isoformat() for i in range(7)}
    hrs = sorted(c["hours"] for c in rel_all)
    return {"spent": round(total, 2), "spent_7d": round(sum(v for d, v in daily.items() if d in week), 2),
            "by_kind": sorted(({"kind": k, "usd": round(v, 2)} for k, v in by_kind.items() if v >= 0.005), key=lambda x: -x["usd"]),
            "by_module": sorted(({"module": k, "usd": round(v, 2)} for k, v in by_module.items() if v >= 0.005), key=lambda x: -x["usd"]),
            "daily": [{"day": d, "usd": round(daily.get(d, 0), 2)} for d in days],
            "released_14d": len(released_14), "in_progress": sum(1 for c in changes if c["status"] == "candidate"),
            "median_h": hrs[len(hrs) // 2] if hrs else None, "max_h": hrs[-1] if hrs else None,
            "changes": [c for c in changes if c["status"] == "candidate" or c in released_14 or c["spent"] > 0][:40]}


_DELIV: dict[str, tuple] = {}


def _delivery_files(cfg: dict) -> list[dict]:
    """The project's delivery workbooks ([results.delivered]), oldest first: {path, mtime, name, cases: {case: status}}.
    Cached until a file changes."""
    spec = (cfg.get("results") or {}).get("delivered") or {}
    if not spec.get("files"):
        return []
    files = sorted({f for g in spec["files"] for f in glob.glob(g, recursive=True) if os.path.isfile(f)}, key=os.path.getmtime)
    sig = tuple((f, os.path.getmtime(f)) for f in files)
    key = cfg.get("project", "")
    if key in _DELIV and _DELIV[key][0] == sig:
        return _DELIV[key][1]
    from openpyxl import load_workbook
    pat = re.compile((cfg.get("results") or {}).get("key_pattern") or r"(.+)")
    idpat = re.compile(cfg.get("case_id_pattern") or r".+")
    name_pat = re.compile(spec.get("name") or r"^(.*?)(?:_\d+)?_\d{8}")
    out = []
    for f in files:
        try:
            wb = load_workbook(f, read_only=True, data_only=True)
        except Exception:                                                       # noqa: BLE001  one bad file is skipped
            continue
        if spec.get("sheet") and spec["sheet"] not in wb.sheetnames:
            continue
        rows = (wb[spec["sheet"]] if spec.get("sheet") else wb.worksheets[0]).iter_rows(values_only=True)
        head = [str(h or "").strip() for h in next(rows, [])]
        if spec.get("value", "Status") not in head:
            continue
        data = list(rows)
        vi = head.index(spec.get("value", "Status"))

        def norm(v):
            m = pat.search(str(v or "").strip())
            return m.group(1) if m else None
        ki = None
        for k in spec.get("keys") or []:                 # the first key column whose values look like case ids
            if k in head:
                i = head.index(k)
                vals = [norm(r[i]) for r in data[:200] if i < len(r)]
                if vals and sum(1 for v in vals if v and idpat.fullmatch(v)) >= 0.9 * len(vals):
                    ki = i
                    break
        wb.close()
        if ki is None:
            continue
        cases = {}
        for r in data:
            c, v = norm(r[ki]), r[vi] if vi < len(r) else None
            if c and v:
                cases[c] = str(v).strip()
        stem = Path(f).stem
        m = name_pat.search(stem)
        out.append({"path": f, "mtime": os.path.getmtime(f), "name": (m.group(1) if m else stem).replace("_", " ").strip(),
                    "cases": cases})
    _DELIV[key] = (sig, out)
    return out


def _delivered(cfg: dict) -> dict[str, str]:
    """case_id -> delivered status; the newest workbook that has a case wins."""
    out: dict[str, str] = {}
    for f in _delivery_files(cfg):
        out.update(f["cases"])
    return out


def _deliveries(cfg: dict) -> list[dict]:
    """One entry per delivered batch: versions of one workbook (the same cases) collapse to the newest, and a workbook
    that only puts other deliveries together (all its cases are in others, spread over two or more) is left out."""
    by_set: dict[frozenset, dict] = {}
    for f in _delivery_files(cfg):                       # oldest first: a later version replaces an earlier one
        by_set[frozenset(f["cases"])] = f
    groups = list(by_set.values())
    keep = []
    for g in groups:
        others = [o for o in groups if o is not g]
        cs = set(g["cases"])
        touched = [o for o in others if cs & set(o["cases"])]
        covered = set().union(*(set(o["cases"]) for o in touched)) if touched else set()
        if len(touched) >= 2 and cs <= covered and all(len(o["cases"]) < len(cs) for o in touched):
            continue                                     # a combined workbook
        keep.append(g)
    return keep


def batch_table(con, rows, ledger: dict, imported: dict) -> list[dict]:
    """One row per production / rework batch, for comparing batches of one project: its own cases, its own result rows
    (a results file shared by several batches belongs to the newest; the older ones say so), reviewers' verdicts,
    spend and running time. Columns a project does not have are absent, not zero."""
    now_iso = _iso(dt.datetime.now())
    by_project: dict[str, list] = {}
    for b in rows:
        by_project.setdefault(b["project"], []).append(b)
    out = []
    for project, bs in by_project.items():
        cfg = _cfg(project)
        spec = cfg.get("results") or {}
        scope_names = (cfg.get("batch") or {}).get("scope") or ["scope"]
        good_vals = set(spec.get("good") or [])
        owner: dict[str, str] = {}                          # results file -> newest batch using it
        for b in sorted(bs, key=lambda b: b["created"]):
            if spec.get("file") and b["work_dir"]:
                owner[str(Path(b["work_dir"]) / spec["file"])] = b["batch_id"]
        values, details, table = {}, {}, []
        dspec = spec.get("delivered") or {}
        delivered = _delivered(cfg)
        dcount: dict[str, int] = {}
        case_sets: dict[str, set] = {}
        meta: dict[str, dict] = {}
        for b in sorted(bs, key=lambda b: b["created"]):
            bid = b["batch_id"]
            cases = {r[0] for r in con.execute("SELECT case_id FROM batch_case WHERE batch_id=?", (bid,))}
            row = {"batch_id": bid, "name": (json.loads(b["aliases"] or "[]") or [b["purpose"][:60]])[0], "type": b["type"],
                   "status": b["status"], "date": (b["created"] or "")[:10], "closed": (b["closed"] or "")[:10],
                   "scope": dict(zip(scope_names, _split_scope(b["scope"], len(scope_names)))), "n": len(cases)}
            f = str(Path(b["work_dir"]) / spec["file"]) if spec.get("file") and b["work_dir"] else None
            if f and owner.get(f) != bid:
                row["results_in"] = next((x["name"] for x in table if x["batch_id"] == owner[f]), None) or \
                    (json.loads(next(x for x in bs if x["batch_id"] == owner[f])["aliases"] or "[]") or [owner[f]])[0]
            res = {} if not f or "results_in" in row else {c: v for c, v in _results(Path(f), spec).items() if c in cases}
            if res:
                vc, dc = {}, {}
                for v, d in res.values():
                    vc[v] = vc.get(v, 0) + 1
                    values[v] = values.get(v, 0) + 1
                    if d and v not in good_vals:
                        dc[d] = dc.get(d, 0) + 1
                        details[d] = details.get(d, 0) + 1
                row.update({"with_result": len(res), "values": vc, "details": dc})
            # reviewers: the latest verdict per (case, reviewer); a case two reviewers judged is consistent or not
            verd: dict[tuple, str] = {}
            last_review = None
            for case, rev, v, at in con.execute("SELECT case_id, reviewer, verdict, at FROM review_event WHERE batch_id=?"
                                                " AND action='verdict' ORDER BY id", (bid,)):
                verd[(case, rev)] = v
                last_review = at
            if verd:
                per_case: dict[str, list] = {}
                for (case, _), v in verd.items():
                    per_case.setdefault(case, []).append(v)
                final = {c: vs[-1] for c, vs in per_case.items()}
                two = {c: vs for c, vs in per_case.items() if len(vs) > 1}
                good_reviewed = [c for c in final if res.get(c, ("",))[0] in good_vals]
                row["review"] = {"date": (last_review or "")[:10], "n": len(final),
                                 "correct": sum(1 for v in final.values() if v == "correct"),
                                 "wrong": sum(1 for v in final.values() if v == "wrong"),
                                 "unsure": sum(1 for v in final.values() if v == "unsure"),
                                 "double": len(two), "agree": sum(1 for vs in two.values() if len(set(vs)) == 1),
                                 "good_reviewed": len(good_reviewed),
                                 "good_wrong": sum(1 for c in good_reviewed if final[c] == "wrong")}
            spent = sum(u for _, _, u in ledger.get(bid, [])) + sum(u for _, _, u in imported.get(bid, []))
            iv = _run_intervals(con, bid, now_iso)
            if delivered:
                dv = {}
                for c in cases:
                    v = delivered.get(c)
                    if v:
                        dv[v] = dv.get(v, 0) + 1
                        dcount[v] = dcount.get(v, 0) + 1
                row["delivered"] = dv
                row["delivered_n"] = sum(dv.values())
                case_sets[bid] = cases
            row["state"] = _part_state(con, b)[0]                         # running / stopped (process gone) / paused / ...
            row.update({"spent": round(spent, 2), "per_case": round(spent / len(cases), 4) if cases else None,
                        "run_h": round(_union_h(iv), 1) if iv else None,
                        "note": b["status_reason"] if b["status"] == "paused" else None,
                        "pause": _pause_word(b["status_reason"]) if b["status"] == "paused" else None})
            meta[bid] = {"cases": cases, "res": res, "iv": iv, "created": b["created"]}
            table.append(row)
        table = _merge_by_delivery(cfg, table, meta, scope_names, good_vals) if dspec.get("files") else table
        for r in table:
            case_sets[r["batch_id"]] = r.pop("_cases", None) or meta.get(r["batch_id"], {}).get("cases", set())
        groups = dspec.get("groups") or {}
        if groups:                                       # the project's own summary fields ([results.delivered.groups])
            to_group = {st: name for name, sts in groups.items() for st in sts}
            grp = lambda v: to_group.get(v, "Other")      # noqa: E731
            for r in table:
                if r.get("delivered") is not None:
                    d = {}
                    for k, n in r["delivered"].items():
                        d[grp(k)] = d.get(grp(k), 0) + n
                    r["delivered"] = d
                if r.get("_status"):
                    r["_status"] = {c: grp(v) for c, v in r["_status"].items()}
            delivered = {c: grp(v) for c, v in delivered.items()}
            dspec = {**dspec, "order": list(groups) + ["Other"],
                     "good": [n for n, sts in groups.items() if sts and set(sts) <= set(dspec.get("good") or [])],
                     "bad": [n for n, sts in groups.items() if sts and set(sts) <= set(dspec.get("bad") or [])]}
        order = list(spec.get("good") or []) + [v for v in sorted(values, key=lambda v: -values[v]) if v not in good_vals]
        dcount = {}
        for r in table:
            for k, n in (r.get("delivered") or {}).items():
                dcount[k] = dcount.get(k, 0) + n
        dorder = [v for v in (dspec.get("order") or []) if v in dcount] + sorted(v for v in dcount if v not in (dspec.get("order") or []))
        # totals over distinct cases (batches overlap: a re-run holds the cases of the run it repeats), per filter
        dtot = {}
        if delivered:
            sites = {r["scope"].get(scope_names[0], "") for r in table}
            types = {r["type"] for r in table}
            for site in [""] + sorted(sites):
                for ty in [""] + sorted(types):
                    cs = set()
                    for r in table:
                        if (not site or r["scope"].get(scope_names[0], "") == site) and (not ty or r["type"] == ty):
                            cs |= case_sets.get(r["batch_id"], set())
                    st: dict[str, int] = {}
                    status_of = {}
                    for r in table:                      # a case's status: the row it is delivered in
                        for c, v in (r.get("_status") or {}).items():
                            status_of[c] = v
                    for c in cs:
                        v = status_of.get(c) or delivered.get(c)
                        if v:
                            st[v] = st.get(v, 0) + 1
                    dtot[f"{site}|{ty}"] = {"n": len(cs), "delivered_n": sum(st.values()), "delivered": st}
        for r in table:
            r.pop("_status", None)
        out.append({"project": project, "scope_names": scope_names, "good": sorted(good_vals),
                    "delivered_order": dorder, "delivered_good": list(dspec.get("good") or []),
                    "delivered_bad": list(dspec.get("bad") or []), "delivered_totals": dtot,
                    "values": [v for v in order if v in values], "details": sorted(details),
                    "rows": table})
    return out


def _merge_by_delivery(cfg: dict, parts: list[dict], meta: dict, scope_names: list[str], good_vals: set) -> list[dict]:
    """Rows = delivered batches. Each reins batch joins the delivery holding most of its cases (at least half): re-runs,
    splits and finishing runs become one row with their spend and running time added up and, per case, the newest
    run's result. Batches not delivered yet make one row per work package."""
    deliveries = _deliveries(cfg)
    home: dict[str, int] = {}
    for p in parts:
        cs = meta[p["batch_id"]]["cases"]
        best = max(range(len(deliveries)), key=lambda i: len(cs & set(deliveries[i]["cases"])), default=None)
        if best is not None and cs and len(cs & set(deliveries[best]["cases"])) >= 0.5 * len(cs):
            home[p["batch_id"]] = best

    def combine(members: list[dict], cases: set) -> dict:
        res: dict[str, tuple] = {}
        for p in sorted(members, key=lambda p: meta[p["batch_id"]]["created"], reverse=True):
            for c, v in meta[p["batch_id"]]["res"].items():
                if c in cases and c not in res:
                    res[c] = v
        vc, dc = {}, {}
        for v, d in res.values():
            vc[v] = vc.get(v, 0) + 1
            if d and v not in good_vals:
                dc[d] = dc.get(d, 0) + 1
        spent = round(sum(p["spent"] for p in members), 2)
        iv = [x for p in members for x in meta[p["batch_id"]]["iv"]]
        scopes = [tuple(p["scope"].get(k, "") for k in scope_names) for p in members]
        sc = max(set(scopes), key=scopes.count) if scopes else tuple("" for _ in scope_names)
        r = {"scope": dict(zip(scope_names, sc)), "n": len(cases), "runs": len(members),
             "type": "rework" if any(p["type"] == "rework" for p in members) else "production",
             "spent": spent, "per_case": round(spent / len(cases), 4) if cases else None,
             "run_h": round(_union_h(iv), 1) if iv else None, "_cases": cases}
        if res:
            r.update({"with_result": len(res), "values": vc, "details": dc})
        return r

    out = []
    for i, d in enumerate(deliveries):
        members = [p for p in parts if home.get(p["batch_id"]) == i]
        cases = set(d["cases"])
        r = combine(members, cases)
        st = {}
        for v in d["cases"].values():
            st[v] = st.get(v, 0) + 1
        r.update({"batch_id": f"delivery:{Path(d['path']).name}", "name": d["name"], "status": "done", "state": "delivered",
                  "date": dt.datetime.fromtimestamp(d["mtime"]).date().isoformat(), "closed": "",
                  "started": min((meta[p["batch_id"]]["created"] for p in members), default="")[:10],
                  "delivered": st, "delivered_n": sum(st.values()), "_status": dict(d["cases"]), "note": None, "pause": None})
        if not members:
            r["scope"] = {k: "" for k in scope_names}
        out.append(r)
    rest: dict[tuple, list] = {}
    for p in parts:
        if p["batch_id"] not in home:
            rest.setdefault(tuple(p["scope"].get(k, "") for k in scope_names), []).append(p)
    delivered = _delivered(cfg)
    order = ("running", "stopped", "paused", "waiting", "done")
    for sc, members in rest.items():
        cases = set().union(*(meta[p["batch_id"]]["cases"] for p in members))
        r = combine(members, cases)
        st = {}
        for c in cases:
            if delivered.get(c):
                st[delivered[c]] = st.get(delivered[c], 0) + 1
        states = {p["state"] for p in members}
        r.update({"batch_id": "work:" + "-".join(sc), "name": " ".join(_split_scope("-".join(sc), 1)).replace("-", " "),
                  "status": "running" if "running" in states else "open", "state": next((x for x in order if x in states), "done"),
                  "date": max(meta[p["batch_id"]]["created"] for p in members)[:10], "closed": "",
                  "started": min(meta[p["batch_id"]]["created"] for p in members)[:10],
                  "delivered": st, "delivered_n": sum(st.values()), "note": None,
                  "pause": next((p["pause"] for p in members if p.get("pause")), None)})
        out.append(r)
    return out


def _split_scope(scope: str, n: int) -> list[str]:
    """'mansfield-low' with scope parts [council, wp] -> ['mansfield', 'low']; extra dashes stay in the last part."""
    parts = scope.split("-", n - 1) if n > 1 else [scope]
    return parts + [""] * (n - len(parts))
